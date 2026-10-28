from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import hashlib, json, sys

sys.stdout.reconfigure(encoding='utf-8')
SOURCE = Path('D:/.湘潭大学学习文档/小论文/初稿-中文.docx')
OUT = Path('D:/GitCode/Rediff/output/doc/初稿-中文_第四章实验补写.docx')
OUT.parent.mkdir(parents=True, exist_ok=True)
original = SOURCE.read_bytes()
doc = Document(SOURCE)
body = doc.element.body
old = list(body)

def paragraph(text, kind='body'):
    if kind == 'heading':
        el = deepcopy(old[65])
        for child in list(el):
            if child.tag != qn('w:pPr'):
                el.remove(child)
        p = Paragraph(el, doc)
        p.add_run(text)
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.keep_together = True
        return el
    el = OxmlElement('w:p')
    p = Paragraph(el, doc)
    p.style = doc.styles['Normal']
    pf = p.paragraph_format
    pf.first_line_indent = Pt(21)
    pf.space_before = Pt(0)
    pf.space_after = Pt(5)
    pf.line_spacing = 1.25
    pf.keep_with_next = False
    pf.keep_together = False
    pf.widow_control = True
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    ind = p._p.get_or_add_pPr().find(qn('w:ind'))
    ind.set(qn('w:firstLineChars'), '0')
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10.5)
    rf = r._r.get_or_add_rPr().rFonts
    rf.set(qn('w:eastAsia'), '宋体')
    if kind == 'caption':
        pf.first_line_indent = Pt(0)
        pf.space_before = Pt(5)
        pf.space_after = Pt(4)
        pf.keep_with_next = True
        pf.keep_together = True
        r.font.size = Pt(10)
    return el

parts = []
prose = []
def add(text, kind='body'):
    parts.append(paragraph(text, kind))
    prose.append(text)

def table(index):
    t = deepcopy(old[index])
    # Keep each table row intact. Preserve all cell contents and figure assets.
    rows = t.findall(qn('w:tr'))
    for row_index, row in enumerate(rows):
        pr = row.find(qn('w:trPr'))
        if pr is None:
            pr = OxmlElement('w:trPr'); row.insert(0, pr)
        if pr.find(qn('w:cantSplit')) is None:
            pr.append(OxmlElement('w:cantSplit'))
        if index == 72:
            for p_el in row.iter(qn('w:p')):
                Paragraph(p_el, doc).paragraph_format.keep_with_next = row_index < len(rows) - 1
    parts.append(t)

parts.append(deepcopy(old[63]))
add('我们在 AWA2、CUB 和 SUN 三个基准数据集上评估所提方法。首先比较完整训练数据下的 ZSL 与 GZSL 性能，随后考察训练样本受限时的表现，并通过消融实验分析两种粒度的关系约束与状态重加权的作用。最后，利用真实与生成特征的可视化，观察不同训练样本比例下的分布特征。')

add('Experimental Settings', 'heading')
add('数据集。AWA2 包含 37,322 张动物图像，共 50 个类别，其中已见类和未见类分别为 40 类和 10 类；CUB 包含 11,788 张鸟类图像，共 200 个类别，划分为 150 个已见类和 50 个未见类；SUN 包含 14,340 张场景图像，共 717 个类别，划分为 645 个已见类和 72 个未见类。三个数据集分别覆盖动物识别、细粒度鸟类识别和场景识别。实验沿用 ZeroDiff 的类别划分，未见类别的真实图像不参与生成模型训练。')
add('评价指标。ZSL 仅在未见类别中进行分类，报告未见类平均每类 Top-1 准确率 T1。GZSL 的候选标签同时包含已见类与未见类，分别报告两类测试样本的平均每类准确率 S 和 U，并采用调和平均值 H = 2SU/(S + U) 衡量二者的综合表现。所有准确率以百分数表示，文中性能差值均以百分点计。')
add('实现细节。本文基于 ZeroDiff 的两阶段生成框架实现。视觉特征和实例对比表征分别采用其交叉熵与 PaCo 表征分支，视觉特征维度为 2,048。AWA2 和 SUN 分别使用 85 维和 102 维类别属性，CUB 使用由文本描述提取的 1,024 维语义表示。模型采用 Adam 优化器，扩散步数设为 4，关系投影维度为 512。关系学习采用类别均衡采样，每批选取 16 个类别、每类 4 个实例，并按类别分配扩散状态。训练时先在真实已见类数据上更新关系投影器，再固定投影器参数，以关系损失与基础生成目标共同优化视觉生成器。测试时根据未见类语义生成特征，并训练分类器完成识别。')

add('Comparison with State-of-the-Art Methods', 'heading')
add('表 1 汇总了本文方法与代表性方法的比较结果，包括 MSDN、PSVMA、ZSLViT、PSVMA+ 和 AENet 等嵌入式方法，以及 CE-GZSL、FREE、RE-GZSL、VADS、GenZSL、RLVC 和 ZeroDiff 等生成式方法，同时列出近期方法 SAGE。我们重点考察相对于基础生成框架 ZeroDiff 的变化，并结合表中其他方法评价所提方法的性能。不同方法采用的视觉骨干与语义表示并不完全相同，因此，跨方法结果主要反映各自完整配置下的识别表现。')
add('表 1 在 AWA2、CUB 和 SUN 上的 ZSL 与 GZSL 性能比较。T1、U、S 和 H 的单位均为 %；“–”表示未提供结果。', 'caption')
table(67)
add('ZSL 结果。本文方法在 AWA2、CUB 和 SUN 上的 T1 分别达到 86.4%、88.2% 和 77.8%，相对于表中的 ZeroDiff 分别提高 4.9、0.7 和 1.4 个百分点。其中，SUN 上的结果略高于 RLVC 的 77.7%，为表中最高值；AWA2 上低于 GenZSL 的 92.2%，CUB 上低于 RLVC 的 90.1%。因此，所提方法在三个数据集上均改善了基础框架的未见类识别性能，但其相对其他方法的优势随数据集而变化。')
add('GZSL 结果。本文方法在三个数据集上的 H 分别为 81.4%、83.0% 和 59.7%，相对于 ZeroDiff 分别提高 1.3、1.4 和 1.0 个百分点。在 CUB 和 SUN 上，本文方法取得表中最高的 H；在 AWA2 上，其 H 高于 ZeroDiff 和 RLVC，但仍低于 GenZSL 与 SAGE。这些结果表明，在所评估的数据集上，引入关系监督能够改善基础框架的综合识别表现，其中 CUB 和 SUN 上的优势更为突出。')
add('进一步观察 U 与 S，AWA2 上二者相对于 ZeroDiff 分别提高 1.5 和 1.2 个百分点。CUB 上 U 从 80.0% 提高至 82.8%，S 则从 83.2% 小幅下降至 83.1%；SUN 上 U 从 58.5% 提高至 61.6%，S 从 58.8% 下降至 57.9%。后两个数据集上的 H 提升主要伴随未见类准确率的提高，而非已见类与未见类准确率的同步增长。')

add('Limited Training Data', 'heading')
add('为考察训练数据受限时的表现，我们在 AWA2 和 SUN 上分别采用 30% 与 10% 的已见类训练样本，将本文方法与 ZeroDiff 进行比较。表 2 报告两种训练比例下的 ZSL T1 和 GZSL H，以考察完整数据条件下观察到的收益能否延续至较少训练样本的情形。')
add('表 2 已见类训练样本受限时的性能比较。30% 和 10% 表示保留的已见类训练样本比例；T1 与 H 的单位均为 %。', 'caption')
table(69)
add('在 AWA2 上，采用 30% 训练样本时，本文方法的 T1 和 H 分别为 84.2% 和 79.1%，较 ZeroDiff 提高 1.2 和 0.9 个百分点；采用 10% 训练样本时，两项指标为 84.1% 和 77.4%，分别提高 1.2 和 1.0 个百分点。在这两个训练比例下，本文方法均保持了相对于基础框架的性能增益。')
add('在 SUN 上，采用 30% 训练样本时，T1 和 H 分别较 ZeroDiff 提高 1.8 和 1.1 个百分点；采用 10% 训练样本时，增益分别为 2.0 和 1.3 个百分点。不过，随着训练样本比例从 30% 降至 10%，本文方法的 T1 仍从 75.3% 降至 60.1%，H 从 52.3% 降至 31.8%。因此，这组结果支持所提方法在已测试的低样本条件下优于 ZeroDiff，同时也表明训练样本减少仍会明显限制 SUN 上的识别性能。')

add('Ablation Study', 'heading')
add('为分析各设计的作用，我们在三个数据集上比较完整模型与五种变体，结果见表 3。w/o Class Alignment 与 w/o Instance Alignment 分别移除生成阶段的跨类语义关系约束和类内实例关系约束；w/o Class & Instance Alignment 同时移除这两项约束。对于状态权重，w/o GSR 对各状态采用等权监督，Shared State Weights 则对两种粒度使用共享的状态权重。')
add('表 3 双粒度关系约束与状态重加权的消融结果。各数据集分别报告 ZSL T1 和 GZSL H，单位均为 %。', 'caption')
table(72)
add('双粒度关系约束。移除跨类约束后，AWA2、CUB 和 SUN 上的 T1 分别下降 2.1、0.3 和 1.7 个百分点，H 分别下降 3.3、2.3 和 2.1 个百分点；移除类内约束后，T1 分别下降 3.0、1.7 和 2.9 个百分点，H 分别下降 4.8、1.5 和 3.4 个百分点。两种约束在三个数据集的两项指标上均带来收益，但其贡献大小并不一致。例如，AWA2 和 SUN 上移除类内约束造成的 H 降幅更大，而 CUB 上移除跨类约束的 H 降幅更大。这一结果与分别利用类别语义和实例表征提供不同粒度监督的设计动机一致。')
add('同时移除两项生成关系约束后，三个数据集上的 T1 分别降至 80.8%、82.5% 和 72.9%，H 分别降至 75.2%、77.7% 和 55.2%。相较该变体，完整模型的 T1 分别提高 5.6、5.7 和 4.9 个百分点，H 分别提高 6.2、5.3 和 4.5 个百分点。结合单项移除结果，联合使用两种粒度的约束优于仅保留其中一种，支持二者在当前框架中的互补作用。该变体仍保留关系空间校准与调整后的训练组织，其配置与表 1 中的 ZeroDiff 不同。')
add('状态重加权。相对于 w/o GSR，完整模型在三个数据集上的 T1 分别提高 2.9、2.1 和 1.9 个百分点，H 分别提高 1.6、2.2 和 2.1 个百分点。相对于 Shared State Weights，完整模型的 T1 分别提高 3.1、1.6 和 0.9 个百分点，H 分别提高 1.5、1.2 和 1.3 个百分点。共享权重相较等权监督并非在所有指标上都有改善，例如其在 AWA2 上的 T1 为 83.3%，略低于等权变体的 83.5%。相比之下，完整模型在这组对照中取得一致增益，支持按粒度分配状态贡献的设计。')

add('Qualitative Analysis', 'heading')
add('为直观观察生成特征的分布，我们在 AWA2 未见类别上对真实特征与生成特征进行 t-SNE 可视化。如图 2 所示，上、下两排分别对应 ZeroDiff 和本文方法，三列分别采用 100%、30% 与 10% 的已见类训练样本。每个类别展示 100 个生成样本及全部真实样本，以颜色区分类别，并分别用空心圆和星形标记真实特征与生成特征。')
table(75)
parts.append(deepcopy(old[76]))
add('两种方法均形成了若干较集中的生成特征簇，但部分不同类别的生成簇彼此邻近，且多个类别的生成特征与对应真实特征仍存在分离。在 100% 与 30% 训练样本的结果中，本文方法的部分相邻生成簇呈现较清楚的边界，例如图中部灰色生成簇与邻近粉色、紫色生成簇的分离更容易辨认。采用 10% 训练样本时，本文方法仍能形成若干可辨认的生成簇，但蓝色与绿色等相邻生成簇仍存在混叠。这些观察显示，部分局部簇结构有所改善，生成特征与真实特征之间的分布差异以及局部类别混叠仍然存在。')
add('各子图采用相同设置独立降维，跨图的绝对位置和簇间距离不能直接比较。图 2 为表 2 的识别结果提供定性补充，其结论限于所展示样本的二维分布，不能直接推及高维关系误差或逐扩散状态的关系保持程度。')

# Retain the existing hyperparameter heading without adding results or analysis.
parts.append(deepcopy(old[77]))
for el in old[63:78]:
    body.remove(el)
for offset, el in enumerate(parts):
    body.insert(63 + offset, el)

xml = etree.tostring(doc.element, xml_declaration=True, encoding='UTF-8', standalone=True)
with ZipFile(SOURCE) as z, ZipFile(OUT, 'w') as out:
    for member in z.infolist():
        out.writestr(member, xml if member.filename == 'word/document.xml' else z.read(member.filename))

check = Document(OUT)
assert len(check.tables) == len(Document(SOURCE).tables) == 4
assert [[c.text for row in t.rows for c in row.cells] for t in check.tables] == [[c.text for row in t.rows for c in row.cells] for t in Document(SOURCE).tables]
with ZipFile(SOURCE) as a, ZipFile(OUT) as b:
    changed = [n for n in a.namelist() if a.read(n) != b.read(n)]
    assert changed == ['word/document.xml'], changed
    ra = etree.fromstring(a.read('word/document.xml')).find(qn('w:body'))
    rb = etree.fromstring(b.read('word/document.xml')).find(qn('w:body'))
    for x,y in zip(list(ra)[:63],list(rb)[:63]):
        assert etree.tostring(x) == etree.tostring(y)
    assert etree.tostring(list(ra)[78]) == etree.tostring(list(rb)[-2])
    assert etree.tostring(list(ra)[79]) == etree.tostring(list(rb)[-1])
assert SOURCE.read_bytes() == original
Path('tmp/chapter4/chapter4_text.txt').write_text('\n\n'.join(prose), encoding='utf-8')
Path('tmp/chapter4/manifest.json').write_text(json.dumps({'output':str(OUT),'source_sha256':hashlib.sha256(original).hexdigest(),'changed_parts':changed,'body_characters':sum(map(len,prose))},ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'output':str(OUT),'characters':sum(map(len,prose)),'changed_parts':changed},ensure_ascii=False))
