import sys, hashlib, json, shutil
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')
root = Path(__file__).parent
source = Path(r'D:\.湘潭大学学习文档\小论文\初稿-中文.docx')
output = root / 'edited.docx'
original = source.read_bytes()
(root / 'original.docx').write_bytes(original)
(root / 'source.sha256').write_text(hashlib.sha256(original).hexdigest())
rows = [
 ['w/o Class Alignment','84.3','78.1','87.9','80.7','76.1','57.6'],
 ['w/o Instance Alignment','83.4','76.6','86.5','81.5','74.9','56.3'],
 ['w/o GSR','83.5','79.8','86.1','80.8','75.9','57.6'],
 ['Shared State Weights','83.3','79.9','86.6','81.8','76.9','58.4'],
 ['w/o Class & Instance Alignment','80.8','75.2','82.5','77.7','72.9','55.2'],
 ['Full Model (Ours)','86.4','81.4','88.2','83.0','77.8','59.7'],
]
d = Document()
t = d.add_table(rows=8, cols=7)
t.alignment = WD_TABLE_ALIGNMENT.CENTER
t.autofit = False
widths = [6.6] + [(14.65-6.6)/6]*6
for col,w in zip(t.columns,widths): col.width = Cm(w)
for row in t.rows:
 for c,w in zip(row.cells,widths): c.width = Cm(w)
t.cell(0,0).merge(t.cell(1,0)).text='Methods'
for start,label in [(1,'AWA2'),(3,'CUB'),(5,'SUN')]:
 t.cell(0,start).merge(t.cell(0,start+1)).text=label
 t.cell(1,start).text='T1'
 t.cell(1,start+1).text='H'
for i,values in enumerate(rows,2):
 for c,text in zip(t.rows[i].cells,values): c.text=text

def border(parent,edge,val='single',size=6):
 e=OxmlElement('w:'+edge)
 for k,v in [('val',val),('sz',str(size)),('color','000000'),('space','0')]: e.set(qn('w:'+k),v)
 parent.append(e)

pr=t._tbl.tblPr
style=pr.find(qn('w:tblStyle'))
if style is not None: pr.remove(style)
pr.find(qn('w:tblW')).set(qn('w:w'),str(round(14.65/2.54*1440)))
pr.find(qn('w:tblW')).set(qn('w:type'),'dxa')
borders=OxmlElement('w:tblBorders')
for edge in ['top','bottom']: border(borders,edge,size=12)
for edge in ['left','right','insideH','insideV']: border(borders,edge,'nil',0)
pr.append(borders)
mar=OxmlElement('w:tblCellMar')
for side,val in [('top',40),('bottom',40),('left',70),('right',70)]:
 e=OxmlElement('w:'+side); e.set(qn('w:w'),str(val)); e.set(qn('w:type'),'dxa');mar.append(e)
pr.append(mar)
seen=set()
for i,row in enumerate(t.rows):
 trpr=row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
 if i<2: trpr.append(OxmlElement('w:tblHeader'))
 for j,c in enumerate(row.cells):
  if c._tc in seen: continue
  seen.add(c._tc)
  c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  cp=c._tc.get_or_add_tcPr()
  shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'FFFFFF');shade.set(qn('w:val'),'clear');cp.append(shade)
  cb=OxmlElement('w:tcBorders')
  if i==1: border(cb,'bottom',size=6)
  if i==0 and j==0: border(cb,'bottom',size=6)
  if i==0 and j>0: border(cb,'bottom',size=4)
  if len(cb): cp.append(cb)
  for p in c.paragraphs:
   p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
   f=p.paragraph_format;f.space_before=Pt(0);f.space_after=Pt(0);f.line_spacing=1.0
   f.left_indent=Pt(0);f.right_indent=Pt(0);f.first_line_indent=Pt(0)
   ind=p._p.get_or_add_pPr().find(qn('w:ind'))
   for name in ['leftChars','rightChars','firstLineChars']: ind.set(qn('w:'+name),'0')
   f.keep_with_next=i<7;f.widow_control=False
   for r in p.runs:
    r.font.name='Times New Roman';r.font.size=Pt(10.5);r.bold=i==7
    rf=r._r.get_or_add_rPr().rFonts
    for key in ['ascii','hAnsi','eastAsia','cs']: rf.set(qn('w:'+key),'Times New Roman')

with ZipFile(source) as z:
 xml=etree.fromstring(z.read('word/document.xml'))
 body=xml.find(qn('w:body'))
 # Give the continuation cell a bottom border across the full header width.
 continuation=t.rows[1]._tr.findall(qn('w:tc'))[0]
 cb=OxmlElement('w:tcBorders');border(cb,'bottom',size=6)
 continuation.get_or_add_tcPr().append(cb)
 target=None
 for p in body.findall(qn('w:p')):
  text=''.join(p.itertext())
  if 'Ablation Study' in text:
   nxt=p.getnext()
   while nxt is not None and nxt.tag==qn('w:p'):
    if not ''.join(nxt.itertext()).strip(): target=nxt;break
    nxt=nxt.getnext()
   break
 if target is None: raise RuntimeError('Ablation insertion point not found')
 previous=target.getprevious()
 pp=previous.find(qn('w:pPr'))
 if pp is None: pp=OxmlElement('w:pPr');previous.insert(0,pp)
 pp.append(OxmlElement('w:keepNext'))
 body.replace(target,deepcopy(t._tbl))
 updated=etree.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
 with ZipFile(output,'w') as out:
  for member in z.infolist():out.writestr(member,updated if member.filename=='word/document.xml' else z.read(member.filename))

check=Document(output)
assert len(check.tables)==2
assert [[c.text for c in r.cells] for r in check.tables[1].rows[2:]]==rows
with ZipFile(source) as a,ZipFile(output) as b:
 changed=[n for n in a.namelist() if a.read(n)!=b.read(n)]
 assert changed==['word/document.xml'],changed
 aroot=etree.fromstring(a.read('word/document.xml'));broot=etree.fromstring(b.read('word/document.xml'))
 abody=aroot.find(qn('w:body'));bbody=broot.find(qn('w:body'))
 assert len(abody)==len(bbody)
 diffs=[i for i,(ae,be) in enumerate(zip(abody,bbody)) if etree.tostring(ae)!=etree.tostring(be)]
 assert len(diffs)==2,diffs
print(json.dumps({'output':str(output),'table_rows':8,'columns':7,'changed_parts':changed,'changed_body_elements':diffs},ensure_ascii=False))
