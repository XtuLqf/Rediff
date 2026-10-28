import sys, hashlib
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
sys.stdout.reconfigure(encoding='utf-8')
root=Path(__file__).parent
source=Path(r'D:\.湘潭大学学习文档\小论文\初稿-中文.docx')
raw=source.read_bytes()
(root/'original.docx').write_bytes(raw)
(root/'source.sha256').write_text(hashlib.sha256(raw).hexdigest())
data=[['ZeroDiff','83.0','78.2','82.9','76.4','73.5','51.2','58.1','30.5'],['Ours','84.2','79.1','84.1','77.4','75.3','52.3','60.1','31.8']]
d=Document();t=d.add_table(rows=5,cols=9);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
widths=[3.05]+[(14.65-3.05)/8]*8
for col,w in zip(t.columns,widths):col.width=Cm(w)
for row in t.rows:
 for c,w in zip(row.cells,widths):c.width=Cm(w)
t.cell(0,0).merge(t.cell(2,0)).text='Method'
for start,name in [(1,'AWA2'),(5,'SUN')]:
 t.cell(0,start).merge(t.cell(0,start+3)).text=name
 for offset,percent in [(0,'30%'),(2,'10%')]:
  cell=t.cell(1,start+offset).merge(t.cell(1,start+offset+1))
  p=cell.paragraphs[0];p.add_run(percent)
  r=p.add_run('D');r.italic=True
  r=p.add_run('tr');r.italic=True;r.font.superscript=True
 for j in range(4):t.cell(2,start+j).text='T1' if j%2==0 else 'H'
for i,values in enumerate(data,3):
 for c,value in zip(t.rows[i].cells,values):c.text=value
def border(parent,edge,size=6,val='single'):
 e=OxmlElement('w:'+edge)
 for k,v in [('val',val),('sz',str(size)),('color','000000'),('space','0')]:e.set(qn('w:'+k),v)
 parent.append(e)
pr=t._tbl.tblPr
pr.find(qn('w:tblW')).set(qn('w:w'),str(round(14.65/2.54*1440)))
pr.find(qn('w:tblW')).set(qn('w:type'),'dxa')
b=OxmlElement('w:tblBorders')
for edge in ['top','bottom']:border(b,edge,12)
for edge in ['left','right','insideH','insideV']:border(b,edge,0,'nil')
pr.append(b)
mar=OxmlElement('w:tblCellMar')
for side,val in [('top',40),('bottom',40),('left',65),('right',65)]:
 e=OxmlElement('w:'+side);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');mar.append(e)
pr.append(mar)
seen=set()
for i,row in enumerate(t.rows):
 rp=row._tr.get_or_add_trPr();rp.append(OxmlElement('w:cantSplit'))
 if i<3:rp.append(OxmlElement('w:tblHeader'))
 for j,c in enumerate(row.cells):
  if c._tc in seen:continue
  seen.add(c._tc);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  for p in c.paragraphs:
   p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j==0 else WD_ALIGN_PARAGRAPH.CENTER
   f=p.paragraph_format;f.space_before=Pt(0);f.space_after=Pt(0);f.line_spacing=1
   f.left_indent=Pt(0);f.right_indent=Pt(0);f.first_line_indent=Pt(0);f.keep_with_next=i<4
   ind=p._p.get_or_add_pPr().find(qn('w:ind'))
   for name in ['leftChars','rightChars','firstLineChars']:ind.set(qn('w:'+name),'0')
   for r in p.runs:
    r.font.name='Times New Roman';r.font.size=Pt(10.5);r.bold=i==4
    for key in ['ascii','hAnsi','eastAsia','cs']:r._r.get_or_add_rPr().rFonts.set(qn('w:'+key),'Times New Roman')
 for j,tc in enumerate(row._tr.findall(qn('w:tc'))):
  if i<3:
   b=OxmlElement('w:tcBorders')
   if i==2:border(b,'bottom',6)
   elif j>0:border(b,'bottom',4)
   if len(b):tc.get_or_add_tcPr().append(b)
with ZipFile(source) as zin:
 xml=etree.fromstring(zin.read('word/document.xml'));body=xml.find(qn('w:body'))
 tables=body.findall(qn('w:tbl'));assert len(tables)==2
 spacer=tables[0].getnext();assert spacer.tag==qn('w:p') and not ''.join(spacer.itertext()).strip()
 spacer.addnext(deepcopy(t._tbl))
 updated=etree.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True)
 with ZipFile(root/'edited.docx','w') as zout:
  for member in zin.infolist():zout.writestr(member,updated if member.filename=='word/document.xml' else zin.read(member.filename))
check=Document(root/'edited.docx');assert len(check.tables)==3
assert [[c.text for c in r.cells] for r in check.tables[1].rows[3:]]==data
print('Prepared 3-level header and 16 verified values.')
