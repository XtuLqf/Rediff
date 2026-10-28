import sys, hashlib, io, json
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')
root=Path(__file__).parent
source=Path(r'D:\.湘潭大学学习文档\小论文\初稿-中文.docx')
raw=source.read_bytes()
(root/'original.docx').write_bytes(raw)
(root/'source.sha256').write_text(hashlib.sha256(raw).hexdigest())
original_doc=Document(source)
caption=('图 2 不同训练样本比例下 AWA2 未见类特征的 t-SNE 可视化。'
         '上、下排分别为 ZeroDiff 和本文方法，各列依次对应 100%、30% 和 10% 的已见类训练样本。'
         '每类展示 100 个生成样本及全部真实样本；颜色表示类别，空心圆和星形分别表示真实与生成特征。'
         '各子图采用相同设置独立降维。')
images=[]
for method in ['zerodiff','rediff']:
 for pct in [100,30,10]:
  matches=list(Path(r'D:\CodexToDo\Build-plot-features\outputs',method).glob(f'{method}_real_vs_generated_{pct}pct_*.png'))
  assert len(matches)==1, matches
  images.append(matches[0])

d=Document();t=d.add_table(rows=5,cols=3);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
for col in t.columns:col.width=Cm(14.65/3)
for row in t.rows:
 for c in row.cells:c.width=Cm(14.65/3);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
for i,pct in enumerate([100,30,10]):t.cell(0,i).text=f'{pct}% training samples'
t.cell(2,0).merge(t.cell(2,2)).text='ZeroDiff'
t.cell(4,0).merge(t.cell(4,2)).text='Ours'
for n,path in enumerate(images):
 row=1 if n<3 else 3
 p=t.cell(row,n%3).paragraphs[0]
 shape=p.add_run().add_picture(str(path),width=Cm(4.77),height=Cm(4.77))
 # Crop in Word without modifying the original PNG or its scatter data.
 fill=shape._inline.find('.//'+qn('pic:blipFill'))
 src=OxmlElement('a:srcRect')
 for k,v in [('l',round(52/1260*100000)),('t',round(138/1380*100000)),('r',round(26/1260*100000)),('b',round(60/1380*100000))]:src.set(k,str(v))
 fill.insert(1,src)
 shape._inline.docPr.set('descr',f'AWA2 {path.parent.name}: {path.name.split("_")[4]} training samples; circles: real; stars: generated.')

def format_p(p,keep=True,size=10.5):
 f=p.paragraph_format;f.space_before=Pt(0);f.space_after=Pt(0);f.line_spacing=1
 f.left_indent=Pt(0);f.right_indent=Pt(0);f.first_line_indent=Pt(0)
 f.keep_with_next=keep;f.keep_together=True;f.widow_control=False
 ind=p._p.get_or_add_pPr().find(qn('w:ind'))
 for name in ['leftChars','rightChars','firstLineChars']:ind.set(qn('w:'+name),'0')
 for r in p.runs:
  r.font.name='Times New Roman';r.font.size=Pt(size)
  rf=r._r.get_or_add_rPr().rFonts
  for key in ['ascii','hAnsi','cs']:rf.set(qn('w:'+key),'Times New Roman')
  rf.set(qn('w:eastAsia'),'宋体')

seen=set()
for i,row in enumerate(t.rows):
 row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
 for c in row.cells:
  if c._tc in seen:continue
  seen.add(c._tc)
  for p in c.paragraphs:
   p.alignment=WD_ALIGN_PARAGRAPH.CENTER;format_p(p,size=9.5 if i==0 else 10.5)
   if i in (2,4):p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(2)
pr=t._tbl.tblPr
pr.find(qn('w:tblW')).set(qn('w:w'),str(round(14.65/2.54*1440)));pr.find(qn('w:tblW')).set(qn('w:type'),'dxa')
b=OxmlElement('w:tblBorders')
for edge in ['top','bottom','left','right','insideH','insideV']:
 e=OxmlElement('w:'+edge);e.set(qn('w:val'),'nil');b.append(e)
pr.append(b)
mar=OxmlElement('w:tblCellMar')
for side in ['top','bottom','left','right']:
 e=OxmlElement('w:'+side);e.set(qn('w:w'),'0');e.set(qn('w:type'),'dxa');mar.append(e)
pr.append(mar)
cp=d.add_paragraph(caption);cp.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY;format_p(cp,False,10)
cp.paragraph_format.space_after=Pt(6);cp.paragraph_format.line_spacing=1.15
bio=io.BytesIO();d.save(bio)

REL_NS='http://schemas.openxmlformats.org/package/2006/relationships'
with ZipFile(source) as zin,ZipFile(bio) as fragment:
 xml=etree.fromstring(zin.read('word/document.xml'));body=xml.find(qn('w:body'))
 rels=etree.fromstring(zin.read('word/_rels/document.xml.rels'))
 next_id=max(int(r.get('Id')[3:]) for r in rels if r.get('Id','').startswith('rId'))+1
 new_media={};new_map={}
 frels=etree.fromstring(fragment.read('word/_rels/document.xml.rels'))
 for rel in frels:
  if rel.get('Type','').endswith('/image'):
   old_id=rel.get('Id');target=rel.get('Target');fname='tsne_'+Path(target).name
   new_id=f'rId{next_id}';next_id+=1;new_map[old_id]=new_id
   nr=etree.SubElement(rels,'{'+REL_NS+'}Relationship')
   nr.set('Id',new_id);nr.set('Type',rel.get('Type'));nr.set('Target','media/'+fname)
   new_media['word/media/'+fname]=fragment.read('word/'+target)
 table_xml=deepcopy(t._tbl)
 for blip in table_xml.findall('.//'+qn('a:blip')):blip.set(qn('r:embed'),new_map[blip.get(qn('r:embed'))])
 next_shape=max([int(e.get('id')) for e in xml.findall('.//'+qn('wp:docPr'))]+[0])+1
 for e in table_xml.findall('.//'+qn('wp:docPr')):
  e.set('id',str(next_shape));e.set('name',f'tSNE panel {next_shape}');next_shape+=1
 target=None;section_heading=None;section_text=None
 for p in body.findall(qn('w:p')):
  text=''.join(p.itertext())
  if 'Analysis of Dual-Granularity Relation Consistency' in text:
   section_heading=p;section_text=p.getnext();n=section_text.getnext()
   if n is not None and n.tag==qn('w:p') and not ''.join(n.itertext()).strip():target=n
   break
 assert target is not None,'Section 4.4 insertion point missing'
 for p in [section_heading,section_text]:
  pp=p.find(qn('w:pPr'))
  if pp is None:pp=OxmlElement('w:pPr');p.insert(0,pp)
  kn=pp.find(qn('w:keepNext'))
  if kn is None:pp.append(OxmlElement('w:keepNext'))
  else:kn.set(qn('w:val'),'1')
  if pp.find(qn('w:keepLines')) is None:pp.append(OxmlElement('w:keepLines'))
 target.addprevious(table_xml);body.replace(target,deepcopy(cp._p))
 types=etree.fromstring(zin.read('[Content_Types].xml'))
 CT='http://schemas.openxmlformats.org/package/2006/content-types'
 if not any(e.get('Extension')=='png' for e in types):
  ct=etree.SubElement(types,'{'+CT+'}Default');ct.set('Extension','png');ct.set('ContentType','image/png')
 changed={'word/document.xml':etree.tostring(xml,xml_declaration=True,encoding='UTF-8',standalone=True),
          'word/_rels/document.xml.rels':etree.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True),
          '[Content_Types].xml':etree.tostring(types,xml_declaration=True,encoding='UTF-8',standalone=True)}
 with ZipFile(root/'edited.docx','w') as zout:
  for member in zin.infolist():zout.writestr(member,changed.get(member.filename,zin.read(member.filename)))
  for name,content in new_media.items():zout.writestr(name,content)

check=Document(root/'edited.docx')
assert len(check.inline_shapes)==7
assert len(check.tables)==4
assert [p.text for p in check.paragraphs if p.text.strip() and p.text!=caption]==[p.text for p in original_doc.paragraphs if p.text.strip()]
for i in range(3):assert [[c.text for c in r.cells] for r in check.tables[i].rows]==[[c.text for c in r.cells] for r in original_doc.tables[i].rows]
(root/'manifest.json').write_text(json.dumps({'caption':caption,'images':[str(p) for p in images]},ensure_ascii=False,indent=2),encoding='utf-8')
print('Verified six original PNGs, source paragraph text and three existing tables preserved.')
